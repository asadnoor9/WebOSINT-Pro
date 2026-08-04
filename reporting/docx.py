"""
WebOSINT Pro - Word (.docx) report generator

Requires python-docx (the optional 'reports' extra). The value is imported
lazily so the core tool works without it installed.
"""

from __future__ import annotations

from datetime import datetime
from typing import Any

from reporting.base import BaseReport
from reporting.enrichment import enrich_results

SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]


class DocxReport(BaseReport):
    extension = "docx"
    format_name = "Word"
    requires_extra = "reports"

    def __init__(self, results: dict[str, Any] | None = None) -> None:
        super().__init__(results)
        self.meta = self.results.get("meta", {})
        self.summary = self.results.get("summary", {})
        self.findings = self.results.get("findings", {})
        self.errors = self.results.get("errors", {})
        self.normalized = self.results.get("normalized_findings", [])

    def save(self, path: str) -> str:
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
        except ImportError as exc:  # pragma: no cover - depends on optional extra
            raise RuntimeError(
                "Word export needs python-docx. Install the extra: pip install '.[reports]'"
            ) from exc

        enrichment = enrich_results(self.results)
        doc = Document()

        title = doc.add_heading("WebOSINT Pro Report", level=0)
        title.runs[0].font.color.rgb = RGBColor(0x1F, 0x2D, 0x3D)

        # ── Overview ──────────────────────────────────────────────────────────
        doc.add_heading("Overview", level=1)
        overview = doc.add_table(rows=0, cols=2)
        overview.style = "Light Grid Accent 1"
        rows = [
            ("Target", self.meta.get("target", "unknown")),
            ("Mode", self.meta.get("mode", "passive")),
            ("Detail level", self.meta.get("detail_level", "standard")),
            ("TLS verified", str(self.meta.get("tls_verified", True))),
            ("Results trusted", str(self.meta.get("results_trusted", True))),
            ("Generated", datetime.now().isoformat(timespec="seconds")),
            ("Modules run", str(len(self.meta.get("modules_run", [])))),
            ("Total findings", str(self.summary.get("total_findings", 0))),
            ("Unique findings", str(self.summary.get("unique_findings", 0))),
        ]
        for label, value in rows:
            cells = overview.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[0].paragraphs[0].runs[0].font.bold = True

        if not self.meta.get("results_trusted", True):
            warn = doc.add_paragraph()
            run = warn.add_run(
                "TLS verification was disabled for this run. All findings are UNTRUSTED."
            )
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)

        # ── Severity summary ──────────────────────────────────────────────────
        doc.add_heading("Severity summary", level=1)
        counts = self.summary.get("severity_counts", self.summary.get("risk_counts", {}))
        sev_table = doc.add_table(rows=1, cols=2)
        sev_table.style = "Light List Accent 1"
        sev_table.rows[0].cells[0].text = "Severity"
        sev_table.rows[0].cells[1].text = "Count"
        for level in SEVERITY_ORDER:
            cells = sev_table.add_row().cells
            cells[0].text = level.title()
            cells[1].text = str(counts.get(level, 0))

        # ── Findings ──────────────────────────────────────────────────────────
        doc.add_heading("Findings", level=1)
        if self.normalized:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            for index, header in enumerate(["Severity", "Title", "Value", "Status", "Sources"]):
                cell = table.rows[0].cells[index]
                cell.text = header
                cell.paragraphs[0].runs[0].font.bold = True
            for finding in self.normalized:
                cells = table.add_row().cells
                cells[0].text = str(finding.get("severity", "")).upper()
                cells[1].text = str(finding.get("title", ""))
                cells[2].text = _flatten(finding.get("value", ""))[:400]
                cells[3].text = str(finding.get("status", ""))
                cells[4].text = "+".join(finding.get("sources", []))
        else:
            for module, data in self.findings.items():
                doc.add_heading(module, level=2)
                for line in _render_lines(data):
                    para = doc.add_paragraph(line, style="List Bullet")
                    para.paragraph_format.left_indent = Pt(12)

        # ── Remediation ───────────────────────────────────────────────────────
        remediation = enrichment.get("remediation_intelligence", [])
        if remediation:
            doc.add_heading("Remediation intelligence", level=1)
            for item in remediation:
                doc.add_heading(item.get("issue_title", "Issue"), level=2)
                for label, key in [
                    ("Severity", "severity"),
                    ("Confidence", "confidence"),
                    ("Why risky", "why_risky"),
                    ("Impact", "impact"),
                    ("Recommended fix", "recommended_fix"),
                ]:
                    value = item.get(key)
                    if value:
                        para = doc.add_paragraph()
                        para.add_run(f"{label}: ").bold = True
                        para.add_run(str(value))

        if self.errors:
            doc.add_heading("Module errors", level=1)
            for module, message in sorted(self.errors.items()):
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(f"{module}: ").bold = True
                para.add_run(str(message))

        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = "WebOSINT Pro - authorized security testing only"

        doc.save(path)
        return path


def _flatten(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        return "; ".join(_flatten(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}={_flatten(v)}" for k, v in sorted(value.items()))
    return str(value)


def _render_lines(data: Any, depth: int = 0) -> list[str]:
    pad = "    " * depth
    out: list[str] = []
    if isinstance(data, dict):
        for key, value in data.items():
            if str(key).startswith("_") or value in (None, "", [], {}):
                continue
            if isinstance(value, (dict, list)):
                out.append(f"{pad}{key}:")
                out.extend(_render_lines(value, depth + 1))
            else:
                out.append(f"{pad}{key}: {str(value)[:400]}")
    elif isinstance(data, list):
        for item in data[:100]:
            out.append(f"{pad}{_flatten(item)[:400]}")
    else:
        out.append(f"{pad}{str(data)[:400]}")
    return out or [f"{pad}No findings"]
